# -*- coding: utf-8 -*-
"""SB様提示用 月別ライン供給計画(27ヶ月版:2027年1月〜2029年3月、基準15・余裕月20)"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

NAVY = "232278"
PALE = "F2F4F6"
HI = "FBEAEB"
WHITE = "FFFFFF"
F = "Meiryo"

months = [f"2027年{m}月" for m in range(1, 13)] + [f"2028年{m}月" for m in range(1, 13)] + [f"2029年{m}月" for m in range(1, 4)]
days = [10, 8, 8, 8, 10, 8, 9, 9, 8, 10, 8, 8] + [10, 8, 8, 10, 8, 8, 10, 8, 9, 9, 8, 10] + [8, 8, 9]  # 実カレンダーの土日数(2028年はうるう年で106日)
BUFFER = {1, 4, 6, 11, 12}
base = [15] * 27
maxl = [20 if int(m.split("年")[1].rstrip("月")) in BUFFER else 15 for m in months]
notes = ["余裕月:最大20ラインまで増枠可" if x == 20 else "" for x in maxl]
for i, m in enumerate(months):
    mm = int(m.split("年")[1].rstrip("月"))
    if mm in (2, 3, 8, 9) and notes[i] == "":
        notes[i] = "当社主催行事との調整月(基準15ライン)"

wb = Workbook()
ws = wb.active
ws.title = "月別ライン供給計画"
thin = Side(style="thin", color="BBBBBB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def put(r, c, v, bold=False, fill=None, color="1A1A1A", size=10.5, align="center", num_fmt=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = Font(name=F, bold=bold, color=color, size=size)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = border
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if num_fmt:
        cell.number_format = num_fmt
    return cell


put(1, 1, "ウルトラマンIP 店頭イベント 月別ライン供給計画(2027年1月〜2029年3月)", bold=True, size=14, align="left")
ws.cell(row=1, column=1).border = Border()
put(2, 1, "円谷プロダクション作成/単価:10万円(税別)/ライン・日(土日開催・当社2名派遣)/現場諸経費2万円/日・交通運搬費・宿泊費・スーツ送料は別途", size=9.5, align="left")
ws.cell(row=2, column=1).border = Border()

headers = ["年月", "開催日数\n(土日)", "基準ライン数", "基準開催数\n(ライン×日)", "基準月額(税別)", "最大ライン数\n(余裕月)", "最大開催数", "最大月額(税別)", "備考"]
for j, h in enumerate(headers, start=1):
    put(4, j, h, bold=True, fill=NAVY, color=WHITE)

r = 5
for i, m in enumerate(months):
    fill = PALE if i % 2 else None
    hi = HI if maxl[i] == 20 else fill
    put(r, 1, m, fill=fill)
    put(r, 2, days[i], fill=fill)
    put(r, 3, base[i], fill=fill)
    put(r, 4, base[i] * days[i], fill=fill, num_fmt="#,##0")
    put(r, 5, base[i] * days[i] * 100000, fill=fill, num_fmt="#,##0")
    put(r, 6, maxl[i], bold=maxl[i] == 20, fill=hi)
    put(r, 7, maxl[i] * days[i], fill=hi, num_fmt="#,##0")
    put(r, 8, maxl[i] * days[i] * 100000, fill=hi, num_fmt="#,##0")
    put(r, 9, notes[i], fill=fill, align="left", size=9.5)
    r += 1

tb = sum(base[i] * days[i] for i in range(27))
tm = sum(maxl[i] * days[i] for i in range(27))
put(r, 1, "27ヶ月 合計", bold=True, fill=NAVY, color=WHITE)
put(r, 2, sum(days), bold=True, fill=NAVY, color=WHITE)
put(r, 3, "15", bold=True, fill=NAVY, color=WHITE)
put(r, 4, tb, bold=True, fill=NAVY, color=WHITE, num_fmt="#,##0")
put(r, 5, tb * 100000, bold=True, fill=NAVY, color=WHITE, num_fmt="#,##0")
put(r, 6, "最大20", bold=True, fill=NAVY, color=WHITE)
put(r, 7, tm, bold=True, fill=NAVY, color=WHITE, num_fmt="#,##0")
put(r, 8, tm * 100000, bold=True, fill=NAVY, color=WHITE, num_fmt="#,##0")
put(r, 9, "", fill=NAVY)
r += 2

conds = [
    "【前提条件】",
    "・基準15ラインを最低開催数保証の対象とします(未実施の場合も基準開催数分を申し受けます)。",
    "・余裕月(1月・4月・6月・11月・12月)は最大20ラインまで増枠が可能です。増枠分は前月1日の開催確定時に確定します。",
    "・開催会場・日程は開催1ヶ月前までに確定。確定後のキャンセル料:1ヶ月前〜14日前30%/13日前〜7日前50%/6日前〜当日100%。",
    "・北関東エリアは2027年1〜3月は実施できません(4月以降は可)。ウルトラセブンは対象外です。",
    "・単価に含むもの:ヒーロー1体・当社2名(担当+スーツアクター)・1日最大5回公演。",
    "・別途:現場諸経費2万円/日、交通運搬費・宿泊費・スーツ送料(基準料金表)、3連休等の3日目は+10万円/日。",
    "・フル運営オプション(運営ディレクター・運営スタッフ3名・MC・音響・諸経費):+22万円/日(会場単位で選択可)。",
]
for c in conds:
    cell = ws.cell(row=r, column=1, value=c)
    cell.font = Font(name=F, size=9.5, bold=c.startswith("【"))
    cell.alignment = Alignment(horizontal="left", vertical="center")
    r += 1

for col, w in zip("ABCDEFGHI", [13, 9, 11, 11, 15, 11, 11, 15, 36]):
    ws.column_dimensions[col].width = w
ws.row_dimensions[4].height = 32
ws.freeze_panes = "A5"

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "SB様提示用_月別ライン供給計画_27ヶ月版.xlsx")
wb.save(OUT)
print("saved:", OUT)
print("days:", sum(days), "base:", tb, tb * 100000, "max:", tm, tm * 100000)
