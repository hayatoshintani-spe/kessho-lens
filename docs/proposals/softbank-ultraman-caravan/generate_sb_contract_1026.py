# -*- coding: utf-8 -*-
"""SB様提示用 契約内容とノベルティ制作の進め方(10/26お打ち合わせ資料)"""
import os

HERE = os.path.dirname(os.path.abspath(__file__))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

INK = RGBColor(0x1A, 0x1A, 0x1A)
SUB = RGBColor(0x5A, 0x62, 0x6B)
RED = RGBColor(0xC0, 0x00, 0x00)
NAVY = RGBColor(0x23, 0x22, 0x78)
SILVER = RGBColor(0x8E, 0x94, 0x9B)
PALE = RGBColor(0xF2, 0xF4, 0xF6)
PALE_RED = RGBColor(0xFB, 0xEA, 0xEB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TSUB_NAVY = RGBColor(0x1F, 0x33, 0x8C)
FONT = "Meiryo"
ASSETS = os.path.join(HERE, "assets")
BAND_IMG = os.path.join(ASSETS, "band_header.png")
COVER_IMG = os.path.join(ASSETS, "band_cover.png")

SW, SH = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]
page_no = [0]


def set_font(run, size, color=INK, bold=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
    if ea is None:
        from lxml import etree
        ea = etree.SubElement(rPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
    ea.set('typeface', FONT)


def add_slide():
    page_no[0] += 1
    return prs.slides.add_slide(BLANK)


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, gap=5):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        for (t, sz, c, b) in (item if isinstance(item, list) else [item]):
            r = p.add_run()
            r.text = t
            set_font(r, sz, c, b)
    return box


def rect(slide, x, y, w, h, fill=None, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    return sp


def footer(slide):
    textbox(slide, Inches(5.17), Inches(7.16), Inches(3.0), Inches(0.28),
            [("© TSUBURAYA PRODUCTIONS", 8, SILVER, False)], align=PP_ALIGN.CENTER)
    textbox(slide, Inches(10.25), Inches(7.1), Inches(2.2), Inches(0.3),
            [("Strictly Confidential", 10, RED, True)], align=PP_ALIGN.RIGHT)
    textbox(slide, Inches(12.6), Inches(7.1), Inches(0.5), Inches(0.3),
            [(str(page_no[0]), 10, SILVER, False)], align=PP_ALIGN.RIGHT)


def header(slide, title, sub=None):
    slide.shapes.add_picture(BAND_IMG, 0, 0, width=SW, height=Inches(0.86))
    textbox(slide, Inches(0.5), Inches(0.08), Inches(11.2), Inches(0.7),
            [(title, 21, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        textbox(slide, Inches(0.55), Inches(0.98), Inches(12.2), Inches(0.4),
                [(sub, 12.5, SUB, False)])
    footer(slide)


def add_table(slide, x, y, w, headers, rows, col_widths=None, font_size=11.5, row_h=0.5):
    n_rows = len(rows) + 1
    tbl = slide.shapes.add_table(n_rows, len(headers), x, y, w, Inches(row_h * n_rows)).table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = int(w * cw / total)
    for j, htxt in enumerate(headers):
        c = tbl.cell(0, j)
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_left = c.margin_right = Inches(0.08)
        p = c.text_frame.paragraphs[0]
        r = p.add_run(); r.text = htxt
        set_font(r, font_size, WHITE, True)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if i % 2 == 1 else PALE
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = c.margin_right = Inches(0.08)
            text, bold, color = val, False, INK
            if isinstance(val, tuple):
                text, bold = val[0], val[1]
                color = val[2] if len(val) > 2 else INK
            for k, line in enumerate(str(text).split("\n")):
                p = c.text_frame.paragraphs[0] if k == 0 else c.text_frame.add_paragraph()
                r = p.add_run(); r.text = line
                set_font(r, font_size, color, bold)
    return tbl


def block(slide, x, y, w, h, title, lines, fill=PALE, title_color=RED):
    b = rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    b.adjustments[0] = 0.06
    items = [(title, 12.5, title_color, True)] + [(l, 11, INK, False) for l in lines]
    textbox(slide, x + Inches(0.25), y + Inches(0.15), w - Inches(0.5), h - Inches(0.3), items, gap=4)


# 1. 表紙 --------------------------------------------------------------------
s = add_slide()
s.shapes.add_picture(COVER_IMG, 0, 0, width=Inches(4.09), height=SH)
textbox(s, Inches(4.5), Inches(1.25), Inches(8.4), Inches(0.5),
        [("ソフトバンク株式会社 御中", 14, SUB, False)], align=PP_ALIGN.CENTER)
textbox(s, Inches(4.5), Inches(2.25), Inches(8.4), Inches(2.0), [
    ("ウルトラソフトバンクフェス", 30, INK, True),
    ("契約内容とノベルティ制作の進め方", 24, INK, True),
], align=PP_ALIGN.CENTER, gap=10)
textbox(s, Inches(4.5), Inches(4.2), Inches(8.4), Inches(0.9), [
    ("お打ち合わせ資料", 13.5, INK, False),
    ("契約期間 2027年1月〜2028年12月(2年間)・基準15ライン", 13.5, INK, False),
], align=PP_ALIGN.CENTER, gap=4)
textbox(s, Inches(4.5), Inches(5.7), Inches(8.4), Inches(1.0), [
    ("2026年10月26日", 15, TSUB_NAVY, True),
    ("(株)円谷プロダクション", 15, TSUB_NAVY, True),
    ("金額はすべて税別", 10.5, SILVER, False),
], align=PP_ALIGN.CENTER, gap=3)
textbox(s, Inches(5.17), Inches(7.16), Inches(3.0), Inches(0.28),
        [("© TSUBURAYA PRODUCTIONS", 8, SILVER, False)], align=PP_ALIGN.CENTER)
textbox(s, Inches(10.25), Inches(7.1), Inches(2.5), Inches(0.3),
        [("Strictly Confidential", 10, RED, True)], align=PP_ALIGN.RIGHT)

# 2. 本日の進め方 --------------------------------------------------------------
s = add_slide()
header(s, "本日の進め方", "契約の構成と主要条件をご確認いただき、ノベルティ・制作物の進め方と締結までの日程を決めます")
add_table(s, Inches(0.7), Inches(1.55), Inches(7.3),
          ["#", "議題", "目安"],
          [
              ["1", "契約の構成(2通)と主要条件のご確認", "25分"],
              ["2", "ノベルティ・制作物の進め方(品目・数量・申請・許諾料)", "25分"],
              ["3", "キービジュアルと素材の提供", "10分"],
              ["4", "締結までの日程と発注・契約の方法", "10分"],
          ],
          col_widths=[0.3, 4.2, 0.8], row_h=0.55, font_size=12)
block(s, Inches(8.3), Inches(1.55), Inches(4.3), Inches(3.0), "本日決めていただきたいこと", [
    "・契約を「イベント実施契約」と「IPライセンス契約」の2通にすること",
    "・発注の形(期間一括のEDI発注)と電子契約での締結",
    "・貴社内の取りまとめ窓口(1部署)",
    "・ノベルティの品目候補と、数量の申告・報告のしかた",
    "・キービジュアル打ち合わせの日程(11月上旬)",
    "・契約書ドラフトのご確認担当者とご確認期間",
])

# 3. 契約の構成 ----------------------------------------------------------------
s = add_slide()
header(s, "1. 契約の構成(2通)", "イベントの実施(役務)とキャラクターの使用(権利)を分け、それぞれ独立した契約として締結します")
add_table(s, Inches(0.7), Inches(1.55), Inches(11.9),
          ["契約", "内容", "金額(固定分)", "期間", "お支払い"],
          [
              [("(A)イベント実施基本契約", True), "店頭イベントへのキャラクターと担当者の派遣。最低開催数保証付き", ("3億1,500万円\n(10万円/日×3,150現場)", True, RED), "2027年1月〜2028年12月", "四半期前払い(ご相談)\n従量分は実績精算"],
              [("(B)IPライセンス契約", True), "告知・イベントキット・訴求ツール・LP/Web・SNSでのIP使用と監修。作成許諾(20%)を含む", ("4,000万円\n(2,000万円/年×2年)", True, RED), "締結日〜2028年12月", "契約時に一括前払い\n作成許諾料は実績精算"],
          ],
          col_widths=[1.8, 3.2, 1.9, 1.7, 1.9], row_h=0.9, font_size=11)
block(s, Inches(0.7), Inches(4.4), Inches(5.8), Inches(2.3), "2通に分ける理由", [
    "・役務の対価(イベント)と権利の対価(IP使用)を分け、両社の社内処理を明確にするため",
    "・それぞれ単独で成立する契約とし、相互に条件を参照しない書き方にします",
    "・排他条件(他の携帯キャリアと店頭イベント契約をしない)はIPライセンス契約に置きます",
])
block(s, Inches(6.8), Inches(4.4), Inches(5.8), Inches(2.3), "契約書に共通で入る一般条項", [
    "・秘密保持(特別価格と最低保証の条件を含む)",
    "・反社会的勢力の排除、契約違反時の解除、損害賠償、準拠法・管轄",
    "・貴社EDIの発注書と本契約の関係は、本契約を優先する旨を記載",
    "・電子契約での締結を想定(貴社のサービスをご教示ください)",
])

# 4. (A)主要条件 ----------------------------------------------------------------
s = add_slide()
header(s, "1. (A)イベント実施基本契約の主要条件", "10月22日にご提示した実施条件を契約条項に落とし込んだものです")
add_table(s, Inches(0.7), Inches(1.5), Inches(11.9),
          ["条項", "内容"],
          [
              [("役務と単価", True), "ヒーロー1体+当社2名(担当・スーツアクター)、1回30分・1日最大5回。100,000円/ライン・日(本件限りの特別価格)"],
              [("ライン数", True), "基準15ライン。余裕月(1・4・6・11・12月)は最大20ラインまで増枠可(前月1日の開催確定時に確定)"],
              [("最低開催数保証", True), "基準15ライン分=3,150現場を保証。第1期(2027年1〜3月)390現場/第2期(2027年4月〜2028年3月)1,560現場/第3期(2028年4〜12月)1,200現場。未実施でも保証分を申し受け、返還はありません"],
              [("精算", True), "各期末に実績を集計。増枠分は追加精算、最低保証を下回った場合も保証額をお支払いいただきます"],
              [("開催確定とキャンセル", True), "会場・日程は開催1ヶ月前までに確定。確定後のキャンセル料:1ヶ月前〜14日前30%/13日前〜7日前50%/6日前〜当日100%"],
              [("従量費用", True), "現場諸経費2万円/日、3日目+10万円/日、MC+3万円/日、フル運営+22万円/日(会場単位)。交通運搬・宿泊・スーツ送料は基準料金表により実費"],
              [("エリアとキャラクター", True), "関信越エリア(埼玉・茨城・群馬・栃木・山梨・長野・新潟)は2027年1〜3月は実施不可。ウルトラセブンは対象外。8ヒーロー(権利調整により一部変更の場合あり)"],
              [("役割分担と安全", True), "貴社が主催者としてイベント保険・JASRAC申請・会場管理・来場者の安全管理を担当。当社は運営マニュアルとレクチャーを提供。実施確認書を会場ごとに取り交わし"],
              [("協議事項", True), "天災・感染症・会場都合による中止の扱い(キャンセル料と最低保証への算入)は契約書で整理します"],
          ],
          col_widths=[1.5, 4.5], row_h=0.53, font_size=10.5)

# 5. (B)主要条件 ----------------------------------------------------------------
s = add_slide()
header(s, "1. (B)IPライセンス契約の主要条件", "キャラクターの使用範囲と、貴社制作物の作成許諾・監修のルールを定めます")
add_table(s, Inches(0.7), Inches(1.5), Inches(11.9),
          ["条項", "内容"],
          [
              [("許諾の内容", True), "8ヒーローを、店頭イベントの告知、イベントキット、店頭訴求ツール、LP/Web、SNSで使用する権利。作成物の監修を含む。地域は日本国内"],
              [("ライセンス料", True), "4,000万円(2,000万円/年×2年)。契約締結時に一括前払い。返還はありません"],
              [("作成許諾ロイヤリティ", True), "貴社制作のイベントキット・ノベルティ・店頭訴求ツールの作成費用(製造原価)の20%。製造数量ベースの最低保証を前半(2027年1月〜2028年3月)と後半(2028年4〜12月)で設定。未達分は期末確定"],
              [("数量の報告", True), "製造数量は四半期ごとに、製造証憑(工場の納品書など)を添えてご報告ください。当社に確認権を置きます"],
              [("申請と監修", True), "品目ごとに仮申請(回答約15営業日)→デザイン→監修(5営業日)→製造。告知物はKV3営業日・Web流用1〜2営業日・チラシ1〜2営業日"],
              [("制作ルール", True), "クレジット「©円谷プロ」を左下に表記。他社キャラクターとの混在不可(貴社キャラクターは可)。ノベルティはお子様向け玩具・菓子類など不可の品目あり。非公式商品の使用・販売・装飾利用は不可"],
              [("権利と終了後の扱い", True), "制作物に含まれるキャラクター部分の権利は当社に帰属。使用はキャンペーン期間中に限る。契約終了後の在庫・告知物の扱いは契約書で定めます"],
              [("排他", True), "契約期間中、当社は国内の他の携帯キャリア3社と店頭イベントの契約を行いません(店頭イベントに限る)"],
          ],
          col_widths=[1.5, 4.5], row_h=0.6, font_size=10.5)

# 6. ノベルティの進め方 ------------------------------------------------------------
s = add_slide()
header(s, "2. ノベルティ・制作物の進め方", "春節前の量産から逆算し、11月上旬までの仮申請と11月中のキービジュアル確定が前提になります")
steps = [("品目の決定", "10月中"), ("仮申請", "11月上旬まで"), ("当社回答", "約15営業日"), ("デザイン・監修", "5営業日"), ("量産", "12月"), ("配布開始", "2027年1月")]
x0, y0, bw, bh, gapw = Inches(0.7), Inches(1.55), Inches(1.75), Inches(0.95), Inches(0.28)
for i, (t, d) in enumerate(steps):
    x = x0 + i * (bw + gapw)
    b = rect(s, x, y0, bw, bh, fill=NAVY if i in (1, 4) else PALE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    b.adjustments[0] = 0.12
    col = WHITE if i in (1, 4) else INK
    textbox(s, x, y0 + Inches(0.12), bw, bh - Inches(0.2), [(t, 12.5, col, True), (d, 10.5, col, False)], align=PP_ALIGN.CENTER, gap=2)
    if i < len(steps) - 1:
        textbox(s, x + bw, y0 + Inches(0.3), gapw, Inches(0.4), [("▶", 11, SILVER, False)], align=PP_ALIGN.CENTER)
add_table(s, Inches(0.7), Inches(2.8), Inches(11.9),
          ["項目", "内容"],
          [
              [("品目候補", True), "タオル、トートバッグ、ステッカー、クリアファイル、缶バッジ、ステーショナリー、修了証(当社既存商品のステッカーを卸価格でご購入いただく案も可)"],
              [("数量の目安", True), "撮影会1開催(土日2日)あたり約400個。2年間・基準で約1,575開催=約63万個。初回量産(2027年1〜3月分)は約195開催=約7.8万個"],
              [("許諾料の計算", True), "製造原価単価 × 製造数量 × 20%。例:単価35円×7.8万個×20%=約55万円(初回量産分)。四半期ごとに製造証憑を添えてご報告"],
              [("仮申請の入力", True), "企画商品名は「ノベルティ ○○」、カテゴリーは貴社の品目に合わせて選択、仮売価は製造原価単価。詳細は「仮申請のご案内」(10/22配布)"],
              [("混在ルール", True), "他社キャラクターとの同一盤面は不可。貴社キャラクター(おとうさん犬など)との組み合わせは個別にご相談"],
          ],
          col_widths=[1.4, 4.6], row_h=0.66, font_size=10.5)

# 7. KVと素材 ----------------------------------------------------------------
s = add_slide()
header(s, "3. キービジュアルと素材の提供", "キャンペーン名称とキービジュアルを11月中に両社で確定します")
add_table(s, Inches(0.7), Inches(1.5), Inches(11.9),
          ["項目", "内容"],
          [
              [("キービジュアル", True), "貴社で制作(当社での受託も可)。11月上旬に方向性の打ち合わせ、11月中に確定。監修はKV3営業日"],
              [("当社提供素材", True), "集合ビジュアル(60周年ロゴ版)、個体素材、ロゴ、撮影会用のキャラクターなし背景素材。データ形式・解像度は別途ご案内"],
              [("素材の使い分け", True), "撮影会の背景バナーはキャラクターなし素材。集合ビジュアル・個体素材・ロゴは会場装飾と告知物に使用可"],
              [("クレジット", True), "「©円谷プロ」を左下に表記(キャラクターに被る場合は位置を調整)。貴社で加工した制作物の監修は5営業日"],
              [("告知", True), "全会場の情報をまとめた告知ページの開設をご検討ください(当社公式サイトから誘導)。撮影後の商談・アンケート案内は告知物と受付で事前にお知らせ"],
          ],
          col_widths=[1.4, 4.6], row_h=0.7, font_size=11)
textbox(s, Inches(0.7), Inches(6.0), Inches(11.9), Inches(0.8), [
    ("・素材は本キャンペーンに限ってご使用ください。使用分のみをお渡しします", 11, INK, False),
    ("・キービジュアルの打ち合わせ候補日を本日いただければ、当社の監修担当(成岡)を含めて設定します", 11, INK, False),
], gap=4)

# 8. 日程 ----------------------------------------------------------------------
s = add_slide()
header(s, "4. 締結までの日程", "10月末の締結に向けて、契約書ドラフトのご確認期間を確保したく、本日中に進め方を決めさせてください")
add_table(s, Inches(0.7), Inches(1.5), Inches(11.9),
          ["時期", "内容", "担当"],
          [
              [("10月26日(本日)", True), "契約構成・主要条件の確認、取りまとめ窓口、発注・契約方法の確認", "両社"],
              [("10月27〜28日", True), "契約書ドラフト2通を電子契約の形式でご提示", "当社"],
              [("10月29〜30日", True), "貴社でのご確認(法務・購買)、修正協議", "貴社"],
              [("10月31日", True), "契約締結(電子契約)。IP使用料の請求書発行", "両社"],
              [("11月上旬", True), "キービジュアル打ち合わせ、ノベルティ品目の仮申請、運営レクチャー日程の調整", "両社"],
              [("11月中", True), "キービジュアル確定、仮申請の回答", "両社"],
              [("12月", True), "デザイン監修完了、量産開始。1月開催分の会場確定(12月1日)、運営レクチャー", "両社"],
              [("2027年1月", True), "開始(基準15ライン)", "両社"],
          ],
          col_widths=[1.6, 4.0, 0.7], row_h=0.5, font_size=11)
block(s, Inches(0.7), Inches(6.1), Inches(11.9), Inches(0.8), "ご確認したいこと", [
    "貴社の購買規程で最低保証分の期間一括発注と四半期前払いが可能か。電子契約サービスの指定の有無。ドラフトのご確認担当者とご確認に必要な日数",
], fill=PALE_RED)

OUT = os.path.join(HERE, "SB様提示用_契約内容とノベルティ制作の進め方_1026.pptx")
prs.save(OUT)
print("saved:", OUT)
